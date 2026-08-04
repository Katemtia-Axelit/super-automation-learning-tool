# -*- coding: utf-8 -*-
import sys, os, shutil, tempfile
sys.stdout.reconfigure(encoding='utf-8')
import json

base = r'D:\Axelit\工作\trae\超级自动化学习工具'
ppt_folder = os.path.join(base, 'data', '讲义')
out_root = os.path.join(base, 'extracted_pics')
tmp_dir = tempfile.mkdtemp(prefix='ppt_conv_')

import pythoncom, win32com.client

def ppt_to_pptx(src_ppt, dst_pptx):
    pythoncom.CoInitialize()
    try:
        p = win32com.client.Dispatch('PowerPoint.Application')
        p.DisplayAlerts = 2
        try:
            presentation = p.Presentations.Open(src_ppt, ReadOnly=True, WithWindow=False)
            presentation.SaveAs(dst_pptx, 24)
            presentation.Close()
            return True
        finally:
            p.Quit()
    finally:
        pythoncom.CoUninitialize()

def extract_all(ppt_path, out_dir):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import hashlib, os
    os.makedirs(out_dir, exist_ok=True)

    prs = Presentation(ppt_path)
    slides_data = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_info = {'slide_num': slide_idx + 1, 'title': '', 'texts': [], 'images': []}
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t and (not slide_info['title'] or len(t) > len(slide_info['title'])):
                        slide_info['title'] = t
                all_text = shape.text_frame.text.strip()
                if all_text:
                    slide_info['texts'].append(all_text[:300])
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                img_bytes = img.blob
                img_ext = img.ext
                img_hash = hashlib.md5(img_bytes).hexdigest()[:8]
                img_size = len(img_bytes)
                out_path = os.path.join(out_dir, f'slide{slide_idx+1:03d}_{img_hash}.{img_ext}')
                with open(out_path, 'wb') as f:
                    f.write(img_bytes)
                slide_info['images'].append({'hash': img_hash, 'ext': img_ext, 'size': img_size, 'file': os.path.basename(out_path)})
                if img_size > 2000:  # Only print non-blank images
                    print(f'  Slide {slide_idx+1} [{slide_info["title"][:40]}] -> {os.path.basename(out_path)} ({img_size//1024}KB)')
        slides_data.append(slide_info)
    return slides_data

# Process 第3章 and 第4章
for prefix in ['第3章', '第4章']:
    fname = None
    for f in sorted(os.listdir(ppt_folder)):
        if f.startswith(prefix) and f.endswith('.ppt') and not f.endswith('.pptx'):
            fname = f; break
    if not fname:
        print(f'NOT FOUND: {prefix}')
        continue

    src_path = os.path.join(ppt_folder, fname)
    key = prefix.replace('第', 'ch')
    tmp_ppt = os.path.join(tmp_dir, f'{key}.ppt')
    tmp_pptx = os.path.join(tmp_dir, f'{key}.pptx')
    out_dir = os.path.join(out_root, key)
    os.makedirs(out_dir, exist_ok=True)

    print(f'\n=== {fname} ===')
    shutil.copy2(src_path, tmp_ppt)
    ppt_to_pptx(tmp_ppt, tmp_pptx)
    slides = extract_all(tmp_pptx, out_dir)

    out_json = os.path.join(base, f'_slides_{key}.json')
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(slides, f, ensure_ascii=False, indent=2)
    print(f'Saved {len(slides)} slides')

print(f'\nDone. Temp: {tmp_dir}')

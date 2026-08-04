import subprocess
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

for pkg in ['python-docx', 'pdfplumber']:
    try:
        __import__(pkg.replace('-', '_'))
    except ImportError:
        subprocess.check_call([sys.executable, '-m', 'pip', 'install', pkg, '-q'])

import docx
import pdfplumber

base = r"D:\Axelit\工作\大一下期末"
out_dir = "data/extracted"
os.makedirs(out_dir, exist_ok=True)

# All exam files to extract
files_to_extract = [
    # 高数试卷
    ("2024年春高等数学AII期末考试试卷(A卷).pdf", "pdf"),
    ("2024年春高等数学AII期末考试试卷(A卷)参考答案与评分标准.doc", "doc"),
    ("2025年春高等数学AII期末考试试卷（A卷） (1).pdf", "pdf"),
    ("2025年春高等数学AII期末考试试卷参考答案与评分标准（A卷）.doc", "doc"),
    ("2023年春季高等数学AII试题（试卷A）.pdf", "pdf"),
    ("2023年春季高等数学AII试题（试卷A）参考答案与评分标准.doc", "doc"),
    # 线代试卷
    ("2020-2021-2线性代数试卷A.docx", "docx"),
    ("2020-2021-2线性代数试卷A参考答案（装订版）.docx", "docx"),
    ("2021-2022-2线性代数试卷A（参考答案）.docx", "docx"),
    ("2021-2022-2线性代数试卷Ａ.docx", "docx"),
    ("2022-2023-2线性代数试卷A（参考答案）.docx", "docx"),
    ("2022-2023-2线性代数试卷A（试题页）.docx", "docx"),
    ("2023-2024-2《线性代数》试卷A.doc", "doc"),
    ("2023-2024-2《线性代数》试卷A参考答案.pdf", "pdf"),
    ("2024-2025-2《线性代数》试卷A.docx", "docx"),
    ("2024-2025-2《线性代数》试卷A参考答案 (1).pdf", "pdf"),
    ("2025年春季大学数学AII线性代数参考答案与评分标准（A卷）.doc", "doc"),
    ("2025年春季大学数学AII线性代数（A卷）.pdf", "pdf"),
]

for fname, ftype in files_to_extract:
    fpath = os.path.join(base, fname)
    if not os.path.exists(fpath):
        print(f"SKIP (not found): {fname}")
        continue
    
    safe_name = fname.replace(' (1)', '').replace(' ', '_')
    out_path = os.path.join(out_dir, safe_name.replace(f'.{ftype}', '.txt'))
    
    try:
        if ftype == "pdf":
            text_parts = []
            with pdfplumber.open(fpath) as pdf:
                for i, page in enumerate(pdf.pages):
                    t = page.extract_text()
                    if t:
                        text_parts.append(f"--- Page {i+1} ---\n{t}")
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(text_parts))
            print(f"OK PDF: {fname} -> {len(text_parts)} pages")
            
        elif ftype == "docx":
            doc = docx.Document(fpath)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(paras))
            print(f"OK DOCX: {fname} -> {len(paras)} paragraphs")
            
        elif ftype == "doc":
            # .doc files need special handling - try as docx first
            try:
                doc = docx.Document(fpath)
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
            except:
                paras = [f"[Cannot read .doc file: {fname}]"]
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(paras))
            print(f"OK DOC: {fname} -> {len(paras)} paragraphs")
            
    except Exception as e:
        print(f"ERROR: {fname}: {e}")

print("\nAll done!")
print(f"\nExtracted files in {out_dir}:")
for f in sorted(os.listdir(out_dir)):
    fpath = os.path.join(out_dir, f)
    size = os.path.getsize(fpath)
    print(f"  {f} ({size} bytes)")

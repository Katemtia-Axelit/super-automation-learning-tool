import subprocess, sys, os
sys.stdout.reconfigure(encoding='utf-8')
for pkg in ['python-docx', 'pdfplumber']:
    try: __import__(pkg.replace('-', '_'))
    except: subprocess.check_call([sys.executable, '-m', 'pip', 'install', pkg, '-q'])

import docx, pdfplumber

base = r"D:\Axelit\工作\大一下期末"
out_dir = "data/extracted/linalg"
os.makedirs(out_dir, exist_ok=True)

files = [
    ("2020-2021-2线性代数试卷A.docx", "docx"),
    ("2020-2021-2线性代数试卷A参考答案（装订版）.docx", "docx"),
    ("2021-2022-2线性代数试卷A.docx", "docx"),
    ("2021-2022-2线性代数试卷A（参考答案）.docx", "docx"),
    ("2022-2023-2线性代数试卷A（试题页）.docx", "docx"),
    ("2022-2023-2线性代数试卷A（参考答案）.docx", "docx"),
    ("2023-2024-2《线性代数》试卷A.doc", "doc"),
    ("2023-2024-2《线性代数》试卷A参考答案.pdf", "pdf"),
    ("2024-2025-2《线性代数》试卷A.docx", "docx"),
    ("2024-2025-2《线性代数》试卷A参考答案 (1).pdf", "pdf"),
]

for fname, ftype in files:
    fpath = os.path.join(base, fname)
    if not os.path.exists(fpath):
        print(f"NOT FOUND: {fname}")
        continue
    safe = fname.replace(' ', '_').replace('（', '(').replace('）', ')')
    out = os.path.join(out_dir, safe.rsplit('.', 1)[0] + '.txt')
    try:
        if ftype == "docx":
            doc = docx.Document(fpath)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
        elif ftype == "pdf":
            parts = []
            with pdfplumber.open(fpath) as pdf:
                for i, pg in enumerate(pdf.pages):
                    t = pg.extract_text()
                    if t: parts.append(f"--- Page {i+1} ---\n{t}")
        elif ftype == "doc":
            try:
                doc = docx.Document(fpath)
                parts = [p.text for p in doc.paragraphs if p.text.strip()]
            except:
                parts = [f"[Cannot read: {fname}]"]
        with open(out, 'w', encoding='utf-8') as f:
            f.write('\n'.join(parts))
        print(f"OK: {fname} -> {len(parts)} items")
    except Exception as e:
        print(f"ERROR: {fname}: {e}")

print("\nAll done!")
